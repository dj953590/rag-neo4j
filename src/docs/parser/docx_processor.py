from abc import ABC

from docx import Document
from src.docs.parser.base_processor import BaseProcessor


class DOCXProcessor(BaseProcessor, ABC):
    def extract_text(self):
        """
        Extract text from a DOCX file.
        """
        doc = Document(self.file_path)
        text_blocks = []
        for paragraph in doc.paragraphs:
            text_blocks.append({
                "type": "text",
                "text": paragraph.text.strip(),
                "bounding_box": [0, 0, 0, 0],  # DOCX does not provide bounding boxes
                "page": 1  # DOCX does not have pages
            })
        return text_blocks

    def extract_tables(self):
        """
        Extract tables from a DOCX file.
        """
        doc = Document(self.file_path)
        tables = []
        for table in doc.tables:
            table_data = []
            for row in table.rows:
                row_data = [cell.text for cell in row.cells]
                table_data.append(row_data)
            tables.append({
                "type": "table",
                "data": table_data,
                "bounding_box": [0, 0, 0, 0],  # DOCX does not provide bounding boxes
                "page": 1  # DOCX does not have pages
            })
        return tables

    def extract_images(self, output_dir="extracted_images"):
        """
        Extract images from a DOCX file and perform OCR.
        """
        doc = Document(self.file_path)
        images = []
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                image_data = rel.target_part.blob
                image_ext = rel.target_ref.split(".")[-1]
                image_path = self.save_images(image_data, output_dir, f"image_{len(images)}.{image_ext}")

                # Perform OCR on the image
                ocr_text = self.perform_ocr(image_path)
                images.append({
                    "type": "image",
                    "path": image_path,
                    "text": ocr_text,
                    "bounding_box": [0, 0, 0, 0],  # DOCX does not provide bounding boxes
                    "page": 1  # DOCX does not have pages
                })
        return images
